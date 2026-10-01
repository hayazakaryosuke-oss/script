#!/usr/bin/env python3
"""mdの脚本 → Word。第20稿.docxのスタイル（Title/Subtitle/Cast/Heading1/Action/Dialogue/Transition）を流用する。
使い方: python3 build_docx.py 入力.md 出力.docx"""
import re, sys, copy
from docx import Document

src, dst = sys.argv[1], sys.argv[2]
tpl = 'anata/あなたはだんだん眠くなる_第20稿.docx'
doc = Document(tpl)
body = doc.element.body
for el in list(body):
    if el.tag.endswith('}p') or el.tag.endswith('}tbl'):
        body.remove(el)

lines = open(src, encoding='utf8').read().split('\n')
title = lines[0].lstrip('# ').strip()
meta = [l.strip() for l in lines[1:10] if l.strip() and not l.startswith('凡例') and not l.startswith('#') and not l.startswith('作者メモ') and not l.startswith('決まり')]
doc.add_paragraph(title, style='Title')
for m in meta:
    doc.add_paragraph(m, style='Subtitle')
i = 0
while i < len(lines) and not lines[i].startswith('## 登場人物'):
    i += 1
i += 1
doc.add_paragraph('登場人物', style='Heading1')
while i < len(lines) and not lines[i].startswith('---'):
    if lines[i].startswith('- '):
        doc.add_paragraph(lines[i][2:].strip(), style='Cast')
    i += 1
for l in lines[i:]:
    s = l.rstrip()
    if not s or s == '---':
        continue
    if s.startswith('### '):
        doc.add_paragraph(s[4:], style='Heading1')
    elif s == '＊':
        doc.add_paragraph('＊', style='Transition')
    elif s == '終わり':
        doc.add_paragraph('終わり', style='Transition')
    elif re.match(r'^[^「」\s。]{1,12}「.*」$', s):
        doc.add_paragraph(s, style='Dialogue')
    else:
        doc.add_paragraph(s.lstrip('　'), style='Action')
doc.save(dst)
print('wrote', dst)
